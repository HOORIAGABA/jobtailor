"""Real bytes in, ordered text out.

Fixtures are built in memory rather than committed as binaries, so the
two-column PDF that motivates `engine.layout` is readable as source: the test
shows the layout it is defending against.
"""
from __future__ import annotations

import io

import pytest

from app.domain.errors import ExtractionEmpty, UnsupportedFileType
from app.io.extract import MAX_BYTES, extract_text, suffix_of


# ── fixtures ──────────────────────────────────────────────────────────

def two_column_pdf() -> bytes:
    """A sidebar resume: skills on the left, experience on the right."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(
        pymupdf.Rect(40, 40, 200, 800),
        "R. KHAN\nr.khan@example.com\n\nSKILLS\n\nLanguages\n"
        "Python, SQL, JavaScript\n\nTools\nAirflow, Docker\n\n"
        "EDUCATION\n\nBSc Computer Science\nPunjab University\n2018 - 2022\n",
        fontsize=9, fontname="helv",
    )
    page.insert_textbox(
        pymupdf.Rect(220, 40, 555, 800),
        "EXPERIENCE\n\nData Analyst\nAcme Corp | Jan 2022 - Present\n"
        "- Worked on data pipelines for the reporting team\n"
        "- Built a dashboard that reduced manual reporting\n\n"
        "Analytics Intern\nBeta Labs | 2021\n"
        "- Wrote SQL queries for the marketing team\n",
        fontsize=9, fontname="helv",
    )
    out = doc.tobytes()
    doc.close()
    return out


def single_column_pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(
        pymupdf.Rect(40, 40, 555, 800),
        "R. KHAN\nr.khan@example.com\n\nEXPERIENCE\n\nData Analyst, Acme Corp\n"
        "- Worked on data pipelines for the reporting team\n"
        "- Built a dashboard that reduced manual reporting\n\n"
        "SKILLS\nPython, SQL, Airflow, Docker\n",
        fontsize=10, fontname="helv",
    )
    out = doc.tobytes()
    doc.close()
    return out


def table_layout_docx() -> bytes:
    """The invisible-table template: a huge share of resumes are built this way."""
    from docx import Document

    document = Document()
    document.add_paragraph("R. KHAN")
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.text = "SKILLS"
    left.add_paragraph("Python, SQL, Airflow")
    right.text = "EXPERIENCE"
    right.add_paragraph("Data Analyst at Acme Corp")
    right.add_paragraph("- Worked on data pipelines for the reporting team")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("BSc Computer Science, Punjab University, 2018 - 2022")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def position(text: str, needle: str) -> int:
    index = text.find(needle)
    assert index >= 0, f"{needle!r} missing from extracted text:\n{text}"
    return index


# ── PDF ───────────────────────────────────────────────────────────────

def test_a_two_column_pdf_is_not_interleaved():
    """The whole reason `engine.layout` exists.

    A naive extractor returns "R. KHAN EXPERIENCE r.khan@example.com Data
    Analyst SKILLS ...". Here the sidebar must be finished before the main
    column starts.
    """
    text = extract_text(two_column_pdf(), "resume.pdf")

    last_sidebar = max(
        position(text, "Python, SQL, JavaScript"),
        position(text, "Punjab University"),
    )
    assert last_sidebar < position(text, "EXPERIENCE")


def test_a_two_column_pdf_keeps_a_job_with_its_bullets():
    text = extract_text(two_column_pdf(), "resume.pdf")
    title = position(text, "Data Analyst")
    bullet = position(text, "Worked on data pipelines")
    intern = position(text, "Analytics Intern")
    assert title < bullet < intern


def test_a_single_column_pdf_reads_straight_down():
    text = extract_text(single_column_pdf(), "resume.pdf")
    assert position(text, "R. KHAN") < position(text, "EXPERIENCE")
    assert position(text, "EXPERIENCE") < position(text, "SKILLS")


def test_nothing_is_lost_from_a_two_column_pdf():
    text = extract_text(two_column_pdf(), "resume.pdf")
    for expected in (
        "r.khan@example.com", "Airflow, Docker", "BSc Computer Science",
        "Acme Corp", "Beta Labs", "Wrote SQL queries for the marketing team",
    ):
        assert expected in text


def test_a_pdf_with_no_text_layer_says_so():
    """A scan has no text to read. That must be an error, not empty output."""
    import pymupdf

    doc = pymupdf.open()
    doc.new_page(width=595, height=842)
    data = doc.tobytes()
    doc.close()

    with pytest.raises(ExtractionEmpty) as excinfo:
        extract_text(data, "scan.pdf")
    assert "scan" in str(excinfo.value).lower()


# ── DOCX ──────────────────────────────────────────────────────────────

def test_a_docx_built_as_a_table_is_read():
    """`document.paragraphs` skips table cells and would return almost nothing.

    Silently, which is what makes it dangerous.
    """
    text = extract_text(table_layout_docx(), "resume.docx")
    assert "Python, SQL, Airflow" in text
    assert "Worked on data pipelines for the reporting team" in text


def test_docx_paragraphs_outside_the_table_survive_in_order():
    text = extract_text(table_layout_docx(), "resume.docx")
    assert position(text, "R. KHAN") < position(text, "SKILLS")
    assert position(text, "SKILLS") < position(text, "EDUCATION")


# ── plain text ────────────────────────────────────────────────────────

SAMPLE_TXT = (
    "R. KHAN\nr.khan@example.com\n\nEXPERIENCE\n"
    "Data Analyst, Acme Corp, Jan 2022 - Present\n"
    "- Worked on data pipelines for the reporting team\n"
    "SKILLS\nPython, SQL, Airflow\n"
)


def test_plain_text_passes_through():
    text = extract_text(SAMPLE_TXT.encode(), "resume.txt")
    assert "Worked on data pipelines" in text


def test_an_odd_encoding_is_decoded_rather_than_crashing():
    text = extract_text(SAMPLE_TXT.encode("cp1252"), "resume.txt")
    assert "Data Analyst" in text


# ── cleanup ───────────────────────────────────────────────────────────

def test_bullet_glyphs_become_one_marker():
    """Nine different bullet characters would look like nine kinds of line."""
    raw = "EXPERIENCE\n• first item here\n▪ second item here\n· third item here\n"
    text = extract_text((raw + SAMPLE_TXT).encode(), "r.txt")
    assert "•" not in text and "▪" not in text
    assert "- first item here" in text


def test_ligatures_are_expanded():
    """`workﬂow` is a single glyph in many PDF fonts and matches nothing."""
    raw = "EXPERIENCE\nBuilt a workﬂow for the ofﬁce reporting team\n"
    text = extract_text((raw + SAMPLE_TXT).encode(), "r.txt")
    assert "workflow" in text and "office" in text


def test_runs_of_blank_lines_collapse():
    raw = "EXPERIENCE\n\n\n\n\nData Analyst\n" + SAMPLE_TXT
    assert "\n\n\n" not in extract_text(raw.encode(), "r.txt")


# ── rejection ─────────────────────────────────────────────────────────

def test_an_unsupported_extension_names_what_is_supported():
    with pytest.raises(UnsupportedFileType) as excinfo:
        extract_text(b"x" * 500, "resume.pages")
    assert ".pdf" in str(excinfo.value)


def test_an_oversized_file_is_refused_before_it_is_parsed():
    with pytest.raises(UnsupportedFileType):
        extract_text(b"x" * (MAX_BYTES + 1), "huge.pdf")


def test_an_empty_upload_is_refused():
    with pytest.raises(ExtractionEmpty):
        extract_text(b"", "resume.pdf")


def test_a_near_empty_document_is_refused_rather_than_guessed_at():
    """Three words is not a resume, and the parser would invent one around it."""
    with pytest.raises(ExtractionEmpty):
        extract_text(b"Hello there\n", "resume.txt")


def test_suffix_is_case_insensitive():
    assert suffix_of("Resume.PDF") == ".pdf"
    assert extract_text(SAMPLE_TXT.encode(), "RESUME.TXT")


# ── artifacts seen on a real resume ───────────────────────────────────
# Everything below was found by running scripts/parse_resume.py --text-only on
# an actual two-page Europass CV, not invented.

def test_markdown_links_keep_the_label_and_drop_the_target():
    """Some readers emit `[label](url)` for a link annotation.

    On a resume the target is a URL-ified copy of the label, so keeping both
    doubles the text and gives the parser two candidates for one field.
    """
    raw = (
        "CONTACT\nLinkedIn: [www.linkedin.com/in/hooria-attas]"
        "(https://www.linkedin.com/in/hooria-attas)\n"
        "Website [www.ist.edu.pk](https://www.ist.edu.pk)\n"
    )
    text = extract_text((raw + SAMPLE_TXT).encode(), "r.txt")
    assert "LinkedIn: www.linkedin.com/in/hooria-attas" in text
    assert "Website www.ist.edu.pk" in text
    assert "](" not in text and "https://" not in text


def test_an_invisible_character_does_not_survive_as_a_blank_line():
    """A zero-width space is not whitespace to Python.

    So `" ".join(line.split())` leaves it as a non-empty line, the blank-run
    collapse never fires, and the extracted resume comes out double-spaced.
    """
    raw = "EXPERIENCE\n\n​\n\nData Analyst\n" + SAMPLE_TXT
    text = extract_text(raw.encode(), "r.txt")
    assert "​" not in text
    assert "\n\n\n" not in text


def test_a_soft_hyphen_is_removed():
    raw = "EXPERIENCE\nBuilt a pipe­line for reporting\n" + SAMPLE_TXT
    assert "pipeline" in extract_text(raw.encode(), "r.txt")


# ── wrapped bullets ───────────────────────────────────────────────────

def test_a_bullet_wrapped_across_lines_becomes_one_line():
    """Three printed lines are one claim.

    Left split, each is a separate chance for the parser to drop half a
    sentence, and parse_check compares line by line — so a correctly parsed
    bullet would look like two missing lines.
    """
    raw = (
        "EXPERIENCE\n"
        "- Designed and deployed a production multi-modal RGB-Thermal\n"
        "object-detection pipeline achieving 96.5% mAP@50, exported to\n"
        "ONNX and deployed on NVIDIA Jetson edge hardware.\n"
        + SAMPLE_TXT
    )
    text = extract_text(raw.encode(), "r.txt")
    joined = [l for l in text.splitlines() if l.startswith("- Designed")]
    assert len(joined) == 1
    assert "NVIDIA Jetson edge hardware." in joined[0]


def test_separate_bullets_are_not_merged():
    raw = (
        "EXPERIENCE\n"
        "- Diagnosed and resolved critical model-architecture bugs that\n"
        "blocked production deployment.\n"
        "- Collaborated with data engineers on reliable data pipelines.\n"
        + SAMPLE_TXT
    )
    bullets = [l for l in extract_text(raw.encode(), "r.txt").splitlines()
               if l.startswith("- Diagnosed") or l.startswith("- Collaborated")]
    assert len(bullets) == 2
    assert bullets[0].endswith("blocked production deployment.")
    assert "Collaborated" not in bullets[0]


def test_a_bullet_does_not_swallow_the_heading_after_it():
    """The stop condition that keeps unwrapping safe."""
    raw = (
        "WORK EXPERIENCE\n"
        "- Wrote SQL queries for the marketing team and produced the\n"
        "weekly report.\n"
        "EDUCATION AND TRAINING\n"
        "BSc Computer Science\n"
        + SAMPLE_TXT
    )
    text = extract_text(raw.encode(), "r.txt")
    assert "EDUCATION AND TRAINING" in text.splitlines()


def test_a_blank_line_ends_a_bullet():
    raw = (
        "EXPERIENCE\n"
        "- Built a dashboard that reduced manual reporting\n"
        "\n"
        "Analytics Intern\n"
        + SAMPLE_TXT
    )
    lines = extract_text(raw.encode(), "r.txt").splitlines()
    assert "Analytics Intern" in lines


def test_numbered_bullets_wrap_too():
    raw = (
        "PROJECTS\n"
        "1. Trained a model to predict customer churn using scikit-learn\n"
        "and deployed it as a nightly scheduled job.\n"
        "2. Wrote the evaluation harness.\n"
        + SAMPLE_TXT
    )
    lines = extract_text(raw.encode(), "r.txt").splitlines()
    assert any(l.startswith("1.") and "nightly scheduled job." in l for l in lines)
    assert any(l.startswith("2.") for l in lines)


def test_prose_without_bullets_is_left_alone():
    """The summary paragraph must not be glued to the heading above it."""
    raw = (
        "ABOUT ME\n"
        "AI/ML Engineer with hands-on experience in building and deploying\n"
        "production machine learning systems.\n"
        + SAMPLE_TXT
    )
    lines = extract_text(raw.encode(), "r.txt").splitlines()
    assert "ABOUT ME" in lines
