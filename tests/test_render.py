"""S10 — the deliverable, and the proof it can be read back.

Every rule asserted here is a measured ATS parsing failure, not a preference.
The tests are named after the failure they prevent, so an edit that breaks one
says which resume it would have broken.
"""
from __future__ import annotations

import pytest

from app.domain.models import Bullet, Contact, Item, ResumeDoc, Section
from app.engine import ats
from app.io import render as render_mod


def _doc() -> ResumeDoc:
    return ResumeDoc(
        contact=Contact(
            full_name="Hooria Attas",
            email="hooria@example.com",
            phone="+92 300 0000000",
            location="Islamabad, Pakistan",
            linkedin="linkedin.com/in/hooria",
        ),
        summary="AI engineer building automations and backend services.",
        sections=[
            Section(id="sec.summary", kind="summary", heading="ABOUT ME", items=[]),
            Section(id="sec.experience", kind="experience", heading="WORK EXPERIENCE",
                    items=[Item(
                        id="exp.1", title="ML Associate", org="NexPred",
                        dates="[ 01/03/2026 - 13/08/2026 ]",
                        date_start="2026-03", date_end="2026-08",
                        bullets=[Bullet(id="exp.1.b.1",
                                        text="Built a 4-channel detection pipeline")])]),
            Section(id="sec.skills", kind="skills", heading="TECHNICAL SKILLS",
                    items=[Item(id="skl.1", title="Backend", bullets=[
                        Bullet(id="skl.1.b.1", text="Python FastAPI, Flask, n8n")])]),
            Section(id="sec.certifications", kind="certifications",
                    heading="CERTIFICATIONS", items=[Item(
                        id="crt.1", title="Associate AI Engineer", org="DataCamp",
                        dates="Feb 2025", date_start="2025-02", date_end="")]),
            Section(id="sec.custom.1", kind="custom", heading="PUBLICATIONS",
                    items=[Item(id="cst.1", title="A paper", org="A journal")]),
        ],
    )


# ══ headings ══════════════════════════════════════════════════════════
# Recognition is keyword matching on the heading, so a heading the candidate
# invented is a section the parser does not classify.

def test_a_custom_heading_becomes_the_standard_one():
    """`ABOUT ME` is a Europass heading. No parser looks for it."""
    text = ats.build(_doc()).plain_text()
    assert "PROFESSIONAL SUMMARY" in text
    assert "ABOUT ME" not in text
    assert "TECHNICAL SKILLS" not in text and "SKILLS" in text


def test_an_unclassified_section_keeps_its_own_heading():
    """`custom` means this system could not classify it.

    Substituting a standard name there would be a guess printed on someone's
    resume, so the verbatim heading survives.
    """
    assert "PUBLICATIONS" in ats.build(_doc()).plain_text()


def test_every_known_kind_has_a_standard_heading():
    from app.domain.models import Kind
    from typing import get_args
    covered = set(ATS := ats.ATS_HEADINGS) | {"custom"}
    assert set(get_args(Kind)) == covered, f"unmapped kinds: {set(get_args(Kind)) - covered}"


# ══ structure ═════════════════════════════════════════════════════════

def test_the_name_is_the_first_line_of_the_body():
    """Most ATS engines ignore header and footer regions entirely.

    A resume whose name lives in a Word header arrives anonymous.
    """
    resume = ats.build(_doc())
    assert resume.lines[0].role == "name"
    assert resume.lines[0].text == "Hooria Attas"
    assert resume.lines[1].role == "contact"


def test_contact_details_are_one_plain_line_not_a_table():
    contact = ats.build(_doc()).lines[1].text
    for part in ("hooria@example.com", "+92 300 0000000", "Islamabad, Pakistan"):
        assert part in contact
    assert contact.count("|") >= 2      # pipe-joined, a separator that survives


def test_title_org_and_dates_stay_on_one_line():
    """A table separates a date from its job title in the parser's view."""
    entry = next(l for l in ats.build(_doc()).lines if l.role == "entry")
    assert "ML Associate" in entry.text
    assert "NexPred" in entry.text
    assert "Mar 2026" in entry.text


def test_the_summary_is_not_printed_twice():
    """The summary section folded into `doc.summary` during normalize."""
    text = ats.build(_doc()).plain_text()
    assert text.count("PROFESSIONAL SUMMARY") == 1


def test_a_skills_group_is_plain_text_on_one_line():
    """Not a table, not a rating bar, not an image — no OCR is ever run."""
    text = ats.build(_doc()).plain_text()
    assert "Backend: Python FastAPI, Flask, n8n" in text


# ══ dates ═════════════════════════════════════════════════════════════
# Mixing "Jan 2021 - Present" with "01/2021 -" extracts wrong.

def test_dates_are_rendered_from_the_iso_fields_not_the_printed_string():
    """The source said `[ 01/03/2026 - 13/08/2026 ]`, brackets and days."""
    entry = next(l for l in ats.build(_doc()).lines if l.role == "entry")
    assert "Mar 2026 - Aug 2026" in entry.text
    assert "01/03/2026" not in entry.text and "[" not in entry.text


def test_a_certification_with_one_date_does_not_say_present():
    """`Feb 2025 - Present` reads as a credential still being earned."""
    text = ats.build(_doc()).plain_text()
    assert "Feb 2025" in text
    assert "Feb 2025 - Present" not in text


def test_an_open_ended_job_does_say_present():
    item = Item(id="exp.9", title="Engineer", org="Acme", date_start="2025-01")
    assert ats.date_line(item, "experience") == "Jan 2025 - Present"
    assert ats.date_line(item, "certifications") == "Jan 2025"


def test_an_unparseable_date_falls_back_to_what_the_candidate_wrote():
    """Printing their own text beats printing nothing."""
    item = Item(id="exp.9", title="X", org="Y", dates="Summer 2024")
    assert ats.date_line(item) == "Summer 2024"


# ══ characters ════════════════════════════════════════════════════════
# Smart quotes, em dashes and ligatures are stripped mid-extraction,
# sometimes taking an adjacent word with them.

@pytest.mark.parametrize("raw, expected", [
    ("“scaled” the team", '"scaled" the team'),
    ("2019 – 2022", "2019 - 2022"),
    ("re‑architected", "re-architected"),          # non-breaking hyphen
    ("ﬁnance", "finance"),                          # fi ligature
    ("• built a thing", "- built a thing"),
    ("✓ shipped", "shipped"),
    ("café loyalty app", "cafe loyalty app"),
    ("soft­hyphen", "softhyphen"),
    ("a​b", "ab"),                                  # zero-width space
])
def test_characters_that_do_not_survive_extraction_are_replaced(raw, expected):
    assert ats.sanitize(raw) == expected


def test_sanitised_text_is_always_ascii():
    for raw in ("“x”", "你好", "naïve — really", "\U0001f600"):
        assert ats.sanitize(raw).isascii()


def test_a_url_broken_across_lines_is_rejoined():
    """Extraction rejoins a wrapped URL with a space.

    The resume then ships an unclickable link and feeds the parser a broken
    token. There is exactly one correct reading of a space inside a URL.
    """
    mangled = ("Link: https://www.datacamp.com/completed/track/ "
               "d507ff3346e623c56? utm_medium=organic_social")
    fixed = ats.close_url_gaps(mangled)
    assert "track/d507ff3346e623c56?utm_medium=organic_social" in fixed
    assert " d507" not in fixed


def test_close_url_gaps_leaves_ordinary_prose_alone():
    prose = "Built a thing and shipped it to production"
    assert ats.close_url_gaps(prose) == prose


# ══ the checklist ═════════════════════════════════════════════════════

def test_a_clean_document_reports_no_problems():
    assert ats.problems(ats.build(_doc())) == []


def test_a_document_with_no_name_is_reported():
    doc = _doc()
    doc.contact.full_name = ""
    found = ats.problems(ats.build(doc))
    assert any("name" in p for p in found)


def test_an_overlong_bullet_is_reported():
    doc = _doc()
    doc.sections[1].items[0].bullets[0].text = "x " * (MAX := MAX_BULLET_CHARS_LOCAL())
    assert any("characters" in p for p in ats.problems(ats.build(doc)))


def MAX_BULLET_CHARS_LOCAL() -> int:
    return ats.MAX_BULLET_CHARS


def test_there_is_no_ats_score_anywhere():
    """A number here would be a target, and the thing producing the document
    is the thing that would optimise it — the `EvidenceIndex` argument again."""
    source = (ats.__file__,)
    import pathlib
    body = pathlib.Path(source[0]).read_text(encoding="utf-8")
    for banned in ("def score", "ats_score", "def rating", "return 100"):
        assert banned not in body


# ══ the files ═════════════════════════════════════════════════════════

def test_docx_and_pdf_both_read_back_complete():
    """The one claim in this project checkable end to end with no model.

    If this system's own extractor cannot recover the text, an ATS has no
    better chance.
    """
    out = render_mod.render(_doc())
    assert out.docx[:2] == b"PK"            # a real zip container
    assert out.pdf[:5] == b"%PDF-"
    assert len(out.checks) == 2
    for check in out.checks:
        assert check.missing == [], f"{check.fmt} lost: {check.missing}"
        assert check.ats_problems == []
    assert out.is_clean


def test_the_pdf_has_a_real_text_layer():
    """An image-based PDF contains no extractable text and no OCR is run."""
    from app.io.extract import extract_text
    out = render_mod.render(_doc(), verify_output=False)
    recovered = extract_text(out.pdf, "resume.pdf")
    assert "Hooria Attas" in recovered
    assert "NexPred" in recovered


def test_the_docx_has_no_tables_and_no_header_content():
    import io as _io
    from docx import Document
    out = render_mod.render(_doc(), verify_output=False)
    document = Document(_io.BytesIO(out.docx))
    assert document.tables == []
    for section in document.sections:
        assert not "".join(p.text for p in section.header.paragraphs).strip()
        assert not "".join(p.text for p in section.footer.paragraphs).strip()


def test_both_formats_say_the_same_thing():
    """One `AtsResume`, two renderings. They cannot disagree about content."""
    from app.engine.parse_check import _squash
    from app.io.extract import extract_text
    out = render_mod.render(_doc(), verify_output=False)
    from_docx = _squash(extract_text(out.docx, "r.docx"))
    from_pdf = _squash(extract_text(out.pdf, "r.pdf"))
    for probe in ("HooriaAttas", "MLAssociate", "NexPred", "PythonFastAPI"):
        assert probe.lower() in from_docx.lower()
        assert probe.lower() in from_pdf.lower()


def test_rendering_an_empty_document_fails_loudly():
    with pytest.raises(render_mod.RenderFailed):
        render_mod.render(ResumeDoc())


def test_the_filename_is_one_a_recruiter_can_file():
    assert render_mod.filename_for(_doc(), "AI Engineer") == \
        "Hooria_Attas_AI_Engineer.docx"
    assert render_mod.filename_for(ResumeDoc(), suffix="pdf") == "resume.pdf"


def test_render_returns_bytes_never_paths():
    """Free-tier filesystems are ephemeral; v1 stored paths to vanished files."""
    out = render_mod.render(_doc(), verify_output=False)
    assert isinstance(out.docx, bytes) and isinstance(out.pdf, bytes)


# ══ the job title ═════════════════════════════════════════════════════

def test_the_job_title_goes_under_the_name():
    """A recruiter scanning a stack learns which opening this is for.

    An ATS also keyword-matches the job title with high weight, and
    `brief.role` is the posting's own phrasing, which is what it matches
    against.
    """
    resume = ats.build(_doc(), role="AI Engineer")
    assert [l.role for l in resume.lines[:3]] == ["name", "title", "contact"]
    assert resume.lines[1].text == "AI Engineer"


def test_a_resume_renders_without_a_role():
    """Rendering must not depend on a brief existing."""
    resume = ats.build(_doc())
    assert [l.role for l in resume.lines[:2]] == ["name", "contact"]
    assert not any(l.role == "title" for l in resume.lines)


def test_the_role_reaches_the_rendered_files():
    out = render_mod.render(_doc(), role="AI Engineer", verify_output=False)
    from app.io.extract import extract_text
    assert "AI Engineer" in extract_text(out.docx, "r.docx")
    assert "AI Engineer" in extract_text(out.pdf, "r.pdf")


# ══ skills ════════════════════════════════════════════════════════════

def test_a_skills_group_prints_the_split_terms_not_the_raw_line():
    """The raw line carried every artefact of the source document.

    Measured: a stray `...etc`, a full stop and a comma mid-list, a nested
    `API Python frameworks: Flask` label, and pipes mixed with commas.
    """
    from app.domain.models import Bullet, Item
    item = Item(id="skl.1", title="Backend & Production Pipelines", bullets=[
        Bullet(id="skl.1.b.1",
               text="Python FastAPI | API Python frameworks: Flask, "
                    "FastAPI - Uvicorn | CI-friendly testing | ...etc"),
        Bullet(id="skl.1.b.2", text="| Github, Github actions"),
    ])
    printed = ats.skill_lines(item)[0].text

    assert printed.startswith("Backend & Production Pipelines: ")
    for term in ("Python FastAPI", "Flask", "Uvicorn", "CI-friendly testing",
                 "Github actions"):
        assert term in printed
    assert "...etc" not in printed
    assert "API Python frameworks:" not in printed
    assert "|" not in printed


def test_what_is_printed_and_what_is_protected_are_one_list():
    """Two lists of the same fact is the bug this project has an invariant on.

    `set_skills` may not drop a skill the resume prints, so the printed list
    and the protected list have to be the same list — and they are, because
    both come from `split_skill_line`.
    """
    from app.domain.models import Bullet, Item, ResumeDoc, Section
    from app.engine.normalize import existing_skills

    doc = ResumeDoc(sections=[Section(
        id="sec.skills", kind="skills", heading="SKILLS", items=[Item(
            id="skl.1", title="Core", bullets=[Bullet(
                id="skl.1.b.1",
                text="YOLO (multi-modal / 4-channel) | Python & libraries "
                     "(Numpy, Pytorch) | FastAPI - Uvicorn")])])])

    printed = ats.skill_lines(doc.section_of_kind("skills").items[0])[0].text
    protected = set(existing_skills(doc))

    for term in printed.split(": ", 1)[1].split(", "):
        assert term.lower() in protected, f"{term} printed but not protected"


def test_a_modifier_in_brackets_is_not_printed_as_a_skill():
    from app.domain.models import Bullet, Item
    item = Item(id="skl.1", title="Vision", bullets=[Bullet(
        id="skl.1.b.1", text="YOLO (multi-modal / 4-channel), Scikit-learn")])
    printed = ats.skill_lines(item)[0].text
    assert "YOLO" in printed and "Scikit-learn" in printed
    assert "multi-modal" not in printed and "4-channel" not in printed


# ══ certifications ════════════════════════════════════════════════════

def test_tracking_parameters_are_stripped_from_a_certificate_url():
    """85 of 180 characters on a real entry were utm_* tokens.

    A recruiter reads the line; an ATS indexes every one of those tokens.
    """
    url = ("https://www.datacamp.com/completed/track/d507ff33"
           "?utm_medium=organic_social&utm_campaign=sharewidget&utm_content=soa")
    assert ats.clean_url(url) == "https://www.datacamp.com/completed/track/d507ff33"


def test_a_query_string_the_page_needs_survives():
    """A shorter URL that 404s is worse than a long one."""
    url = "https://example.com/verify?id=abc123&utm_source=email"
    assert ats.clean_url(url) == "https://example.com/verify?id=abc123"


def test_the_link_label_is_dropped():
    assert ats.tidy_link_text("Link: https://example.com/x") == "https://example.com/x"


def test_a_credential_id_is_split_out_of_the_issuer():
    """The parse glued two facts into the organisation field.

    The issuer is what a recruiter scans for and what an ATS maps to a known
    certification authority, so it cannot read as a sentence.
    """
    assert ats.split_credential("BCG X Credential ID: 9ungT9MMB8YDwizqx") == \
        ("BCG X", "Credential ID: 9ungT9MMB8YDwizqx")
    assert ats.split_credential("DataCamp") == ("DataCamp", "")


def test_a_certification_entry_reads_cleanly_end_to_end():
    from app.domain.models import Bullet, Item, ResumeDoc, Section
    doc = ResumeDoc(sections=[Section(
        id="sec.certifications", kind="certifications", heading="CERTS",
        items=[Item(
            id="crt.1", title="GenAI Job Simulation",
            org="BCG X Credential ID: 9ungT9MMB8YDwizqx",
            date_start="2025-02",
            bullets=[Bullet(id="crt.1.b.1",
                            text="Link: https://forage.com/c/x?utm_source=share")])])])
    text = ats.build(doc).plain_text()
    assert "GenAI Job Simulation | BCG X | Feb 2025" in text
    assert "Credential ID: 9ungT9MMB8YDwizqx" in text
    assert "https://forage.com/c/x" in text
    assert "utm_source" not in text and "Link:" not in text


def test_a_bare_link_is_not_bulleted_as_an_achievement():
    from app.domain.models import Bullet, Item
    item = Item(id="crt.1", title="Cert", org="Issuer", bullets=[
        Bullet(id="crt.1.b.1", text="Link: https://example.com/x")])
    roles = [l.role for l in ats.entry_lines(item, "certifications")]
    assert "bullet" not in roles
    assert roles.count("meta") == 1
