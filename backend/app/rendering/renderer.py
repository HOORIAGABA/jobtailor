"""Module 3.4 from the blueprint. Builds an ATS-safe DOCX resume
programmatically from a tailored resume JSON, following professional
ATS-friendly structure: centered header (name, title, contact), then
Summary → Skills → Experience → Projects → Certifications → Education →
Leadership, using real Word bullet lists and plain-text headings."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "data", "outputs")


def contact_info_from_user(user, title: str = "") -> dict:
    """Build the resume header contact_info dict from a user's profile.
    Falls back to the email local-part as the display name so the header
    is never empty."""
    return {
        "full_name": (getattr(user, "full_name", None) or (user.email or "").split("@")[0]).strip(),
        "email": user.email,
        "phone": getattr(user, "phone", None) or "",
        "location": getattr(user, "location", None) or "",
        "linkedin": getattr(user, "linkedin", None) or "",
        "github": getattr(user, "github", None) or "",
        "title": title,
    }

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x55, 0x55)
NAME_SIZE = Pt(18)
TITLE_SIZE = Pt(12)
SECTION_HEADING_SIZE = Pt(12)
BODY_SIZE = Pt(10.5)
META_SIZE = Pt(9.5)

SEPARATOR = " | "

_REPLACEMENT_CHARS = {0xFFFD: " ", 0xFFFE: " ", 0xFFFF: " "}


def _clean(text) -> str:
    if not isinstance(text, str):
        return ""
    return " ".join(text.translate(_REPLACEMENT_CHARS).split()).strip()


def _strip_company_dup(title: str, company: str | None) -> str:
    """Remove a trailing '- Company' / 'Company Company' duplication artifact."""
    t = _clean(title)
    c = _clean(company) if company else ""
    if c and t:
        if t.lower().endswith(c.lower()) and len(t) > len(c):
            t = t[: -len(c)].strip().rstrip(" -–—|").strip()
            t = " ".join(t.split())
    return t


def _add_section_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    run = p.add_run(_clean(text).upper())
    run.bold = True
    run.font.size = SECTION_HEADING_SIZE
    run.font.name = "Calibri"
    run.font.color.rgb = ACCENT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def _add_entry_line(doc: Document, title: str, company: str | None = None, dates: str | None = None):
    p = doc.add_paragraph()
    title = _strip_company_dup(title, company)
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.size = BODY_SIZE
    title_run.font.name = "Calibri"
    if company:
        sep = p.add_run("  -  " + _clean(company))
        sep.font.size = BODY_SIZE
        sep.font.name = "Calibri"
    if dates:
        dates_run = p.add_run("   (" + _clean(dates) + ")")
        dates_run.font.size = BODY_SIZE
        dates_run.font.name = "Calibri"
        dates_run.italic = True
        dates_run.font.color.rgb = MUTED
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    return p


def _add_bullets(doc: Document, bullets: list[str]):
    if not bullets:
        return
    try:
        bullet_style = doc.styles["List Bullet"]
        for b in bullets:
            if b.strip():
                p = doc.add_paragraph(style=bullet_style)
                run = p.add_run(_clean(b))
                run.font.size = BODY_SIZE
                run.font.name = "Calibri"
                p.paragraph_format.space_after = Pt(1)
    except KeyError:
        for b in bullets:
            if b.strip():
                p = doc.add_paragraph()
                run = p.add_run("•  " + _clean(b))
                run.font.size = BODY_SIZE
                run.font.name = "Calibri"
                p.paragraph_format.space_after = Pt(1)


def _add_centered(doc: Document, text: str, size: Pt, bold: bool = False, color: RGBColor | None = None,
                  space_after: Pt = Pt(2)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_clean(text))
    run.bold = bold
    run.font.size = size
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_after = space_after
    return p


def render_resume(tailored_json: dict, output_filename: str, contact_info: dict | None = None) -> str:
    """
    tailored_json: {full_name, summary_heading, summary, skills, experience, projects,
                    education, certifications, leadership, extra_sections, ...}
    contact_info: optional {full_name, title, email, phone, location, linkedin, github}
    Returns the absolute path to the generated .docx file.
    """
    contact_info = contact_info or {}
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(4)

    # =====================================================================
    # HEADER — name, title headline, contact line (plain text, ATS-safe)
    # =====================================================================
    full_name = contact_info.get("full_name") or tailored_json.get("full_name") or "Your Name"
    _add_centered(doc, full_name, NAME_SIZE, bold=True, space_after=Pt(1))

    title = contact_info.get("title")
    if title:
        _add_centered(doc, title, TITLE_SIZE, bold=False, color=MUTED, space_after=Pt(4))

    contact_parts = []
    for key in ("location", "phone", "email", "linkedin", "github"):
        val = _clean(contact_info.get(key))
        if val:
            contact_parts.append(val)
    if contact_parts:
        _add_centered(doc, SEPARATOR.join(contact_parts), META_SIZE, space_after=Pt(6))

    # =====================================================================
    # 1. PROFESSIONAL SUMMARY
    # =====================================================================
    summary = (tailored_json.get("summary") or "").strip()
    if summary:
        _add_section_heading(doc, tailored_json.get("summary_heading") or "Professional Summary")
        sp = doc.add_paragraph(_clean(summary))
        for run in sp.runs:
            run.font.size = BODY_SIZE
            run.font.name = "Calibri"

    # =====================================================================
    # 2. SKILLS
    # =====================================================================
    skills = tailored_json.get("skills") or []
    if skills:
        _add_section_heading(doc, "Skills")
        sp = doc.add_paragraph(", ".join(_clean(s) for s in skills if _clean(s)))
        for run in sp.runs:
            run.font.size = BODY_SIZE
            run.font.name = "Calibri"

    # =====================================================================
    # 3. WORK EXPERIENCE
    # =====================================================================
    experience = tailored_json.get("experience") or []
    if experience:
        _add_section_heading(doc, "Work Experience")
        for exp in experience:
            _add_entry_line(doc, exp.get("title", ""), exp.get("company"), exp.get("dates"))
            _add_bullets(doc, exp.get("bullets", []))

    # =====================================================================
    # 4. PROJECTS
    # =====================================================================
    projects = tailored_json.get("projects") or []
    if projects:
        _add_section_heading(doc, "Projects")
        for proj in projects:
            _add_entry_line(doc, proj.get("title", ""), proj.get("company"), proj.get("dates"))
            _add_bullets(doc, proj.get("bullets", []))

    # =====================================================================
    # 5. CERTIFICATIONS
    # =====================================================================
    certifications = tailored_json.get("certifications") or []
    if certifications:
        _add_section_heading(doc, "Certifications")
        for cert in certifications:
            _add_entry_line(doc, cert.get("title", ""), cert.get("company"), cert.get("dates"))
            _add_bullets(doc, cert.get("bullets", []))

    # =====================================================================
    # 6. EDUCATION
    # =====================================================================
    education = tailored_json.get("education") or []
    if education:
        _add_section_heading(doc, "Education")
        for edu in education:
            _add_entry_line(doc, edu.get("title", ""), edu.get("company"), edu.get("dates"))
            _add_bullets(doc, edu.get("bullets", []))

    # =====================================================================
    # 7. LEADERSHIP / COMMUNITY
    # =====================================================================
    leadership = tailored_json.get("leadership") or []
    if leadership:
        _add_section_heading(doc, "Leadership & Community")
        for lead in leadership:
            _add_entry_line(doc, lead.get("title", ""), lead.get("company"), lead.get("dates"))
            _add_bullets(doc, lead.get("bullets", []))

    # =====================================================================
    # 8. EXTRA SECTIONS (publications, languages, volunteering, etc.)
    # =====================================================================
    extra_sections = tailored_json.get("extra_sections") or []
    for section in extra_sections:
        if not isinstance(section, dict):
            continue
        heading = section.get("heading", "")
        entries = section.get("entries", [])
        if heading and entries:
            _add_section_heading(doc, heading)
            for entry in entries:
                _add_entry_line(doc, entry.get("title", ""), entry.get("company"), entry.get("dates"))
                _add_bullets(doc, entry.get("bullets", []))

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, output_filename)
    doc.save(output_path)
    return output_path