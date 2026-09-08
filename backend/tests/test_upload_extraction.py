"""Regression tests for resume upload text extraction (M4.2).

The old DOCX extractor only read body paragraphs (`document.paragraphs`) so ANY
content inside a table cell - common in designed/skill-matrix resumes - never
reached the LLM. That is what produced "parsed projects have no titles": the
title lines lived in table cells and were silently dropped before the parser
even saw them.
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

from app.routers.resumes import _extract_docx_content


def _build_resume_docx() -> bytes:
    doc = Document()
    doc.add_heading("HOORIA ATTAS", level=0)
    doc.add_paragraph("City: Lahore | Country: Pakistan | Phone: +92 3036800002")
    doc.add_heading("SKILLS", level=1)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Computer Vision\nYOLO\nPyTorch"
    t.cell(0, 1).text = "Backend\nFastAPI\nDocker"
    doc.add_heading("PROJECTS", level=1)
    t2 = doc.add_table(rows=1, cols=1)
    t2.cell(0, 0).text = "FBR Tax Assistant - WhatsApp RAG Chatbot"
    doc.add_paragraph("- Developed a WhatsApp chatbot using RAG")
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph("NexPred Solutions")
    doc.add_paragraph("Machine Learning Associate | [01/03/2026 - Current]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extraction_includes_table_content():
    text = _extract_docx_content(_build_resume_docx())
    assert "FBR Tax Assistant" in text          # project title was in a table cell
    assert "PyTorch" in text and "FastAPI" in text
    assert "HOORIA ATTAS" in text
    assert "Machine Learning Associate" in text


def test_docx_extraction_is_reading_order():
    text = _extract_docx_content(_build_resume_docx())
    assert text.index("SKILLS") < text.index("PROJECTS") < text.index("WORK EXPERIENCE")
    # table cell content appears inside its section (title before the bullet)
    assert text.index("FBR Tax Assistant") < text.index("- Developed a WhatsApp")


if __name__ == "__main__":
    import traceback

    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    fails = 0
    for fn in tests:
        try:
            fn()
            print(f"PASS  {fn.__name__}")
        except Exception:
            fails += 1
            print(f"FAIL  {fn.__name__}")
            traceback.print_exc()
    print(f"{len(tests) - fails}/{len(tests)} passed")
    raise SystemExit(1 if fails else 0)