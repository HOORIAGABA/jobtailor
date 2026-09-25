"""Stage S10 — the deliverable files, and proof that they can be read back.

`engine.ats` decided what the page says. This module only puts those lines into
a .docx and a .pdf, and then reads its own output back to check the decision
survived the file format.

**Returns `bytes`, never paths.** Free-tier hosts have ephemeral filesystems;
v1 stored absolute paths to files that stopped existing on the next restart.
The caller decides whether bytes become a download, an email attachment or a
file on disk.

**Two formats, on purpose.** `.docx` is the more reliably parsed of the two and
is what goes to an ATS. The `.pdf` is for the human: it looks the same on every
machine, which a .docx does not. They are built from one `AtsResume`, so they
cannot say different things.

**Verification, not assertion.** `verify()` renders, then reads the file back
through `io.extract` — the same reader the pipeline uses on an uploaded resume —
and compares what comes out against what went in. This is the one claim in the
project that can be checked end to end without a model, a network or a human:
if this system's own extractor cannot recover the text, an ATS has no better
chance. It is also a closed loop with `engine.layout`, which exists because
two-column resumes defeat naive extraction; the renderer is forbidden from
producing the shape the reader had to be built to survive, and the check is
what keeps that honest.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass

from app.domain.errors import JobTailorError
from app.domain.models import ResumeDoc
from app.engine.ats import AtsResume, Line, build, problems, sanitize
from app.engine.parse_check import _squash
from app.io.extract import extract_text

logger = logging.getLogger(__name__)

# On the ATS-approved list and present on effectively every machine, which is
# the actual requirement: a font the parser's host does not have is a font
# whose ligatures fail. Calibri for Word, Helvetica for PDF because it is one
# of fpdf2's core fonts and needs no font file shipped.
DOCX_FONT = "Calibri"
PDF_FONT = "helvetica"

# A hyphen, not a bullet glyph. `-` is an approved ATS bullet and it is ASCII,
# so it encodes in fpdf2's core-font latin-1 without shipping a Unicode TTF.
# `engine.ats.sanitize` has already turned every decorative bullet into this.
BULLET = "-"

PAGE_MARGIN_MM = 15.0
PDF_BODY_PT = 10.0
PDF_LINE_MM = 4.6


class RenderFailed(JobTailorError):
    """The file could not be produced, or could not be read back."""


# ── docx ──────────────────────────────────────────────────────────────

def to_docx(resume: AtsResume) -> bytes:
    """One column, no tables, no text boxes, no header, no images.

    Every one of those absences is a parsing failure not happening. The
    document is a flat sequence of paragraphs, which is exactly what a parser
    that reads top-to-bottom expects to find.

    Bullets are literal `- ` text rather than Word's list styles. A list style
    puts the bullet in numbering metadata instead of the paragraph text, and
    what survives extraction is the text.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:                              # pragma: no cover
        raise RenderFailed(
            "python-docx is not installed: `pip install python-docx`"
        ) from exc

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(PDF_BODY_PT + 0.5)
    normal.paragraph_format.space_after = Pt(2)
    # Justified text inserts variable word spacing, which extraction reads as
    # runs of whitespace and sometimes as a column break.
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for line in resume.lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(
            f"{BULLET} {line.text}" if line.role == "bullet" else line.text)
        run.font.name = DOCX_FONT

        if line.role == "name":
            run.font.size, run.bold = Pt(18), True
        elif line.role == "title":
            run.font.size, run.bold = Pt(12.5), True
        elif line.role == "meta":
            run.font.size = Pt(9)
        elif line.role == "heading":
            run.font.size, run.bold = Pt(12), True
            paragraph.paragraph_format.space_before = Pt(10)
        elif line.role == "entry":
            run.bold = True
        elif line.role == "contact":
            run.font.size = Pt(10)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── pdf ───────────────────────────────────────────────────────────────

def to_pdf(resume: AtsResume) -> bytes:
    """A real text layer, single column, selectable.

    An image-based PDF is the failure this cannot have: fpdf2 writes text
    operators, so `extract_text` finds words rather than pixels. The round-trip
    check in `verify()` proves it rather than assuming it.
    """
    try:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos
    except ImportError as exc:                              # pragma: no cover
        raise RenderFailed("fpdf2 is not installed: `pip install fpdf2`") from exc

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=PAGE_MARGIN_MM)
    pdf.set_margins(PAGE_MARGIN_MM, PAGE_MARGIN_MM, PAGE_MARGIN_MM)
    pdf.add_page()

    for line in resume.lines:
        if line.role == "name":
            pdf.set_font(PDF_FONT, "B", 17)
        elif line.role == "title":
            pdf.set_font(PDF_FONT, "B", 12)
        elif line.role == "meta":
            pdf.set_font(PDF_FONT, "", 8.5)
        elif line.role == "heading":
            pdf.ln(2.5)
            pdf.set_font(PDF_FONT, "B", 11.5)
        elif line.role == "entry":
            pdf.set_font(PDF_FONT, "B", PDF_BODY_PT)
        elif line.role == "contact":
            pdf.set_font(PDF_FONT, "", 9)
        else:
            pdf.set_font(PDF_FONT, "", PDF_BODY_PT)

        text = f"{BULLET} {line.text}" if line.role == "bullet" else line.text
        # multi_cell wraps within the one column. No table, no cell grid.
        # `new_x=LMARGIN` is load-bearing: without it the cursor stays at the
        # right margin after a wrapped line and the next call has no width to
        # render into, which fpdf2 reports as "not enough horizontal space".
        pdf.set_x(PAGE_MARGIN_MM)
        pdf.multi_cell(
            w=0, h=PDF_LINE_MM, text=text,
            new_x=XPos.LMARGIN, new_y=YPos.NEXT,
        )

    out = pdf.output()
    return bytes(out)


# ── verification ──────────────────────────────────────────────────────

@dataclass
class RenderCheck:
    """What survived the round trip. A checklist, never a score."""
    fmt: str
    byte_count: int
    missing: list[str]
    ats_problems: list[str]

    @property
    def is_clean(self) -> bool:
        return not self.missing and not self.ats_problems

    def summary(self) -> str:
        if self.is_clean:
            return f"{self.fmt}: {self.byte_count} bytes, reads back complete"
        parts = []
        if self.missing:
            parts.append(f"{len(self.missing)} line(s) unreadable")
        if self.ats_problems:
            parts.append(f"{len(self.ats_problems)} ATS problem(s)")
        return f"{self.fmt}: " + ", ".join(parts)


# A line this short is punctuation or a fragment; demanding it survive
# extraction exactly would fail on whitespace handling, not on content.
MIN_LINE_CHARS_TO_CHECK = 4


def verify(resume: AtsResume, data: bytes, filename: str) -> RenderCheck:
    """Read the rendered file back and report what did not survive.

    Comparison is on squashed text — punctuation and spacing removed — because
    a PDF reflows lines and a .docx does not, so matching raw lines would
    report a wrapping difference as lost content. `engine.parse_check._squash`
    is reused deliberately: the same notion of "the same text" that decides
    whether a parse lost anything decides it here too.
    """
    recovered = _squash(extract_text(data, filename))

    missing = [
        line.text for line in resume.lines
        if len(line.text) >= MIN_LINE_CHARS_TO_CHECK
        and _squash(line.text) not in recovered
    ]
    return RenderCheck(
        fmt=filename.rsplit(".", 1)[-1],
        byte_count=len(data),
        missing=missing,
        ats_problems=problems(resume),
    )


@dataclass
class Rendered:
    """The deliverables plus the evidence they are readable."""
    docx: bytes
    pdf: bytes
    text: str
    checks: list[RenderCheck]

    @property
    def is_clean(self) -> bool:
        return all(check.is_clean for check in self.checks)


def render(doc: ResumeDoc, *, role: str = "",
           verify_output: bool = True) -> Rendered:
    """S10. `ResumeDoc` -> .docx + .pdf + the plain text an ATS would see.

    `role` is the job title from the brief, printed under the name.

    `verify_output=False` skips reading the files back. It exists for the case
    where the caller has already checked an identical document, not as a way to
    avoid a failing check — an unverified render is a claim, and this project
    has spent enough time on claims that turned out to be measuring the wrong
    thing.
    """
    resume = build(doc, role)
    if not resume.lines:
        raise RenderFailed(
            "nothing to render: the document has no contact, summary or sections"
        )

    docx_bytes = to_docx(resume)
    pdf_bytes = to_pdf(resume)

    checks: list[RenderCheck] = []
    if verify_output:
        for data, name in ((docx_bytes, "resume.docx"), (pdf_bytes, "resume.pdf")):
            check = verify(resume, data, name)
            checks.append(check)
            if check.missing:
                logger.warning(
                    "%s: %d line(s) did not survive the round trip; first: %r",
                    name, len(check.missing), check.missing[0][:80],
                )
            for problem in check.ats_problems:
                logger.warning("ATS: %s", problem)

    return Rendered(
        docx=docx_bytes, pdf=pdf_bytes,
        text=resume.plain_text(), checks=checks,
    )


def filename_for(doc: ResumeDoc, role: str = "", suffix: str = "docx") -> str:
    """`Hooria_Attas_AI_Engineer.docx` — a name a recruiter can file.

    A resume arriving as `resume(3).pdf` is a small avoidable rudeness, and
    some ATS ingestion paths read the filename as a hint.
    """
    parts = [sanitize(doc.contact.full_name), sanitize(role)]
    stem = "_".join(
        "_".join(p.split()) for p in parts if p
    ) or "resume"
    safe = "".join(c for c in stem if c.isalnum() or c in "_-")
    return f"{safe}.{suffix}"
