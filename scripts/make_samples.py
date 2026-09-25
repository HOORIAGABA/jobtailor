"""Generate synthetic resumes for `samples/`. No real personal data.

    python scripts/make_samples.py

Three layouts, each chosen because it breaks something:

* `two_column.pdf`  — a sidebar. Every naive extractor interleaves it, and the
  parser then sees a skills line wedged between a job title and its bullets.
* `single_column.pdf` — the easy case, as a control. If this one ever fails,
  the bug is not in the layout code.
* `table_layout.docx` — an invisible two-column table, which is how a large
  share of Word resume templates are built. `document.paragraphs` returns
  almost nothing for these, and returns it without error.

These are committed, so a fresh clone can run the pipeline immediately without
anybody's real CV. Put your own resumes in `samples/` too — that folder ignores
everything except these.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

SIDEBAR = """A. MORGAN
a.morgan@example.com
+1 555 0142
Berlin, Germany

SKILLS

Languages
Python, SQL, Go

Frameworks
PyTorch, FastAPI

Tools
Airflow, Docker, Git

EDUCATION

BSc Computer Science
Example University
2017 - 2021
"""

MAIN = """EXPERIENCE

Data Engineer
Northwind Analytics | 03/2022 - Present
- Rebuilt the nightly ingestion pipeline so a failed source retries
  independently, cutting manual reruns from nine a week to none
- Migrated twelve scheduled jobs onto Airflow with explicit dependencies
- Ran the on-call rotation for the data platform

Analytics Intern
Contoso Labs | 06/2021 - 12/2021
- Wrote SQL models for the marketing attribution dashboard
- Helped migrate reporting from spreadsheets to the warehouse

PROJECTS

Churn predictor | 2023
- Trained a gradient-boosted model on two years of subscription data
  and served it behind a FastAPI endpoint
- Wrote the evaluation harness that compared it against the existing rules
"""

SINGLE = f"""A. MORGAN
a.morgan@example.com | +1 555 0142 | Berlin, Germany

SUMMARY

Data engineer with four years building ingestion and reporting pipelines,
most recently owning the nightly platform at a mid-sized analytics company.

{MAIN}
TECHNICAL SKILLS
Languages: Python, SQL, Go
Frameworks: PyTorch, FastAPI
Tools: Airflow, Docker, Git

EDUCATION

BSc Computer Science, Example University, 2017 - 2021
"""


def write_two_column(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(40, 40, 200, 800), SIDEBAR,
                        fontsize=9, fontname="helv")
    page.insert_textbox(pymupdf.Rect(220, 40, 555, 800), MAIN,
                        fontsize=9, fontname="helv")
    doc.save(path)
    doc.close()


def write_single_column(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_textbox(pymupdf.Rect(40, 40, 555, 800), SINGLE,
                        fontsize=10, fontname="helv")
    doc.save(path)
    doc.close()


def write_table_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("A. MORGAN")
    document.add_paragraph("a.morgan@example.com | +1 555 0142 | Berlin, Germany")

    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.text = "SKILLS"
    for line in ("Languages: Python, SQL, Go",
                 "Frameworks: PyTorch, FastAPI",
                 "Tools: Airflow, Docker, Git"):
        left.add_paragraph(line)

    right.text = "EXPERIENCE"
    right.add_paragraph("Data Engineer — Northwind Analytics, 03/2022 - Present")
    for line in (
        "Rebuilt the nightly ingestion pipeline so a failed source retries "
        "independently, cutting manual reruns from nine a week to none",
        "Migrated twelve scheduled jobs onto Airflow with explicit dependencies",
    ):
        right.add_paragraph(f"- {line}")

    document.add_paragraph("EDUCATION")
    document.add_paragraph("BSc Computer Science, Example University, 2017 - 2021")
    document.save(path)


def main() -> int:
    SAMPLES.mkdir(exist_ok=True)
    built = []
    for name, writer in (
        ("two_column.pdf", write_two_column),
        ("single_column.pdf", write_single_column),
        ("table_layout.docx", write_table_docx),
    ):
        path = SAMPLES / name
        writer(path)
        built.append(f"  {path}  ({path.stat().st_size / 1024:.0f} KB)")

    print("Wrote:")
    print("\n".join(built))
    print("\nTry:\n  python scripts/parse_resume.py samples/ --text-only")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
