"""Write the demo résumé as a .docx and the demo posting as a .txt.

The seeded demo (`scripts/demo_seed.py`) writes rows straight into the
database. This writes the two FILES a person would upload and paste, so the
same case can be run through the real UI — upload, confirm, tailor, approve,
send — with `scripts/fake_model.py` answering for the model. See DEMO.md.

    python -m scripts.demo_files            # -> demo/Priya_Raman_CV.docx, demo/posting.txt

The .docx is plain paragraphs on purpose: the point is that its text matches
the fixture exactly, so the coverage check after parsing is clean and every
id the scripted plan cites resolves.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

from scripts.demo_seed import POSTING, RESUME


def write_resume(path: Path) -> None:
    doc = Document()
    contact = RESUME.contact
    doc.add_paragraph(contact.full_name)
    doc.add_paragraph(f"{contact.email} · {contact.phone} · {contact.location}")
    doc.add_paragraph(f"{contact.linkedin} · {contact.github}")
    doc.add_paragraph("SUMMARY")
    doc.add_paragraph(RESUME.summary)
    for section in RESUME.sections:
        doc.add_paragraph(section.heading)
        for entry in section.entries:
            head = " — ".join(x for x in (entry.title, entry.org, entry.dates) if x)
            if head:
                doc.add_paragraph(head)
            for bullet in entry.bullets:
                doc.add_paragraph(bullet, style="List Bullet")
    doc.save(path)


def main(argv: list[str] | None = None) -> int:
    out = Path((argv or sys.argv[1:] or ["demo"])[0])
    out.mkdir(parents=True, exist_ok=True)
    write_resume(out / "Priya_Raman_CV.docx")
    (out / "posting.txt").write_text(POSTING, encoding="utf-8")
    print(f"wrote {out / 'Priya_Raman_CV.docx'}")
    print(f"wrote {out / 'posting.txt'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
